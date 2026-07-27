"""Campaign analysis helpers for intake: read every resource, probe videos, harvest
URLs, classify reference images, and LLM-extract structured rules.

Text extraction degrades gracefully — if an optional parser (pypdf/python-docx/
openpyxl) is missing, the resource is reported as partially handled rather than
silently skipped. Nothing here guesses a rule; ambiguities are surfaced to intake.
"""
import json
import os
import re
import subprocess
import sys
import urllib.request
from math import gcd
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common as C

URL_RE = re.compile(r"https?://[^\s)\"'<>\]}]+")
PLAINTEXT_EXTS = {".txt", ".md", ".markdown", ".log", ".csv", ".tsv"}


# --- URL handling --------------------------------------------------------------
def harvest_urls(text):
    seen, out = set(), []
    for u in URL_RE.findall(text or ""):
        u = u.rstrip(".,);]}'\"")
        if u not in seen:
            seen.add(u)
            out.append(u)
    return out


def classify_url(u):
    low = u.lower()
    if "drive.google.com" in low and ("/folders/" in low or "folderview" in low):
        return "drive_folder"
    if "docs.google.com" in low or "drive.google.com" in low:
        return "gdoc"                # a Google Doc/Sheet/Slides or a Drive file
    if any(h in low for h in ("kick.com", "youtube.com", "youtu.be", "twitch.tv")):
        return "vod"
    return "other"


def _gdoc_export_url(u):
    """Public export-to-text URL for a Google Doc/Sheet, or None."""
    m = re.search(r"/document/d/([A-Za-z0-9_-]+)", u)
    if m:
        return f"https://docs.google.com/document/d/{m.group(1)}/export?format=txt"
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", u)
    if m:
        return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv"
    m = re.search(r"[?&]id=([A-Za-z0-9_-]+)", u)
    if m:
        return f"https://drive.google.com/uc?export=download&id={m.group(1)}"
    return None


def _fetch_text(url, timeout=30):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read().decode("utf-8", errors="replace")


def fetch_gdoc_text(url):
    """Best-effort export of a link-shared Google Doc/Sheet to text. (text, method)."""
    exp = _gdoc_export_url(url)
    if not exp:
        return None, "no export url"
    try:
        txt = _fetch_text(exp)
        if "<html" in txt[:200].lower():     # got a login/redirect page, not the doc
            return None, "not public (login required)"
        return txt, "gdoc-export"
    except Exception as e:
        return None, f"export failed: {e}"


# --- file text extraction ------------------------------------------------------
def extract_text(path):
    """(text, method) or (None, reason). Handlers per file type; optional parsers
    degrade gracefully."""
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext in PLAINTEXT_EXTS:
            return path.read_text(encoding="utf-8", errors="replace"), "text"
        if ext in {".gdoc", ".gsheet", ".gslides"}:
            return _read_gpointer(path)
        if ext == ".json":
            return _read_json_maybe_pointer(path)
        if ext == ".pdf":
            return _read_pdf(path)
        if ext == ".docx":
            return _read_docx(path)
        if ext in {".xlsx", ".xls"}:
            return _read_xlsx(path)
        if ext == ".rtf":
            return _strip_rtf(path.read_text(encoding="utf-8", errors="replace")), "rtf"
    except Exception as e:
        return None, f"error: {e}"
    return None, "no handler for this file type"


def _read_gpointer(path):
    data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    url = data.get("url") or data.get("doc_id") or ""
    if data.get("doc_id") and "http" not in url:
        url = f"https://docs.google.com/document/d/{data['doc_id']}/edit"
    txt, method = fetch_gdoc_text(url)
    if txt:
        return txt, method
    return f"[Google file pointer -> {url}]", "gpointer-url-only"


def _read_json_maybe_pointer(path):
    raw = path.read_text(encoding="utf-8", errors="replace")
    try:
        data = json.loads(raw)
        if isinstance(data, dict) and ("doc_id" in data or "url" in data):
            return _read_gpointer(path)
    except Exception:
        pass
    return raw, "json-text"


def _read_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        return None, "pypdf not installed (pip install pypdf)"
    reader = PdfReader(str(path))
    return "\n".join((pg.extract_text() or "") for pg in reader.pages), "pdf"


def _read_docx(path):
    try:
        import docx
    except ImportError:
        return None, "python-docx not installed (pip install python-docx)"
    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs), "docx"


def _read_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        return None, "openpyxl not installed (pip install openpyxl)"
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(", ".join(cells))
    return "\n".join(rows), "xlsx"


def _strip_rtf(rtf):
    return re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", rtf)


# --- video probe ---------------------------------------------------------------
def probe_video(path):
    """{duration, width, height, aspect} via ffprobe; Nones + note if unavailable."""
    import shutil
    if shutil.which("ffprobe") is None:
        return {"duration_sec": None, "width": None, "height": None,
                "aspect": None, "note": "ffprobe not found"}
    proc = subprocess.run(
        ["ffprobe", "-v", "error", "-select_streams", "v:0",
         "-show_entries", "stream=width,height:format=duration",
         "-of", "json", str(path)],
        stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    try:
        data = json.loads(proc.stdout or "{}")
        st = (data.get("streams") or [{}])[0]
        w, h = st.get("width"), st.get("height")
        dur = float(data.get("format", {}).get("duration") or 0) or None
        aspect = None
        if w and h:
            g = gcd(w, h) or 1
            aspect = f"{w // g}:{h // g}"
        return {"duration_sec": round(dur, 2) if dur else None,
                "width": w, "height": h, "aspect": aspect}
    except Exception as e:
        return {"duration_sec": None, "width": None, "height": None,
                "aspect": None, "note": f"probe error: {e}"}


# --- reference images ----------------------------------------------------------
def reference_purpose(name):
    low = name.lower()
    if "safezone" in low or "safe zone" in low or "safe-zone" in low or "safezones" in low:
        return "safezone"
    if "placement" in low or "example" in low:
        return "placement"
    return None


def image_dims(path):
    try:
        from PIL import Image
        with Image.open(path) as im:
            return im.size
    except Exception:
        return (None, None)


# --- LLM structured extraction -------------------------------------------------
LLM_FIELDS = ("required_elements", "banned_words", "banned_topics", "platform_rules",
              "format_specs", "hashtags", "mentions", "submission_process", "deadlines",
              "payout_terms", "style_guidance", "examples_good", "examples_bad")


def groq_extract(client, campaign, corpus, floor):
    """Structured extraction into the LLM_FIELDS. Returns a dict (may be partial).
    The deterministic `floor` is passed so the model AUGMENTS, never removes."""
    if client is None:
        return {}
    corpus = corpus[:24000]     # keep within context; note truncation upstream
    system = (
        "You are a meticulous campaign-rules analyst for a short-form clipping team. "
        "Extract EVERYTHING relevant from the campaign corpus into STRICT JSON with keys: "
        + ", ".join(LLM_FIELDS) + ". "
        "banned_words = individual words/phrases that must never appear. "
        "banned_topics = themes to avoid. format_specs = {length, aspect_ratio, safe_zones, "
        "resolution} if stated. required_elements = list of {type, detail}. examples_good/"
        "examples_bad = short strings. Use [] or \"\" when unknown — NEVER invent a rule. "
        "Return ONLY the JSON object."
    )
    user = (f"Campaign: {campaign}\n"
            f"Deterministic floor already known (augment, do not drop): {json.dumps(floor)}\n\n"
            f"CORPUS:\n{corpus}")
    raw = C.groq_chat(client, system, user, temperature=0.2, max_tokens=2000)
    m = re.search(r"\{.*\}", raw, re.S)
    if not m:
        C.warn("Groq extraction returned no JSON — keeping deterministic rules only.")
        return {}
    try:
        return json.loads(m.group(0))
    except Exception as e:
        C.warn(f"Groq extraction JSON parse failed ({e}) — deterministic rules only.")
        return {}


def merge_rules(floor, llm):
    """Union lists (floor is a floor; LLM adds, never removes). Scalars: keep floor,
    fill blanks from LLM."""
    out = dict(floor)
    for key in LLM_FIELDS:
        fv, lv = out.get(key), llm.get(key)
        if isinstance(fv, list) or isinstance(lv, list):
            merged, seen = [], set()
            for item in (fv or []) + (lv or []):
                k = json.dumps(item, sort_keys=True) if isinstance(item, (dict, list)) else str(item).strip().lower()
                if k and k not in seen:
                    seen.add(k)
                    merged.append(item)
            out[key] = merged
        else:
            out[key] = fv if fv else (lv or fv)
    return out
